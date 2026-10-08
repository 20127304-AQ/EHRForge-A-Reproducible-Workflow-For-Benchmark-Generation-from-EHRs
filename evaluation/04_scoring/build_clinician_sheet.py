"""PART 5: clinician validation sheet (protocol v1.4 scoring.judge_validation). Run from project root.
Draws ~100 judge-called rows from the FULL judge run, stratified by judge score (1..5) x reader family (extractive/qwen),
10 per cell (=20 per score level); if a cell has < 10 rows, the shortfall is taken from the other family at the same score
level, then (if still short) left short. rng = numpy.random.default_rng(42), cells in fixed order (score 1..5, extractive, qwen).
Writes W/results/clinician/clinician_validation_100.xlsx (blinded) and clinician_validation_100_key.csv (hidden key)."""
import os
import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

import os as _os, sys as _sys  # release adaptation (paths/config only)
_HERE = _os.path.dirname(_os.path.abspath(__file__))
_sys.path.insert(0, _os.path.dirname(_HERE))  # evaluation/ (release_config.py)
import release_config as _rc  # noqa: E402  reads evaluation/config.yaml
_rc.enter_workspace()  # inputs resolve relative to WORKSPACE_DIR (original project layout)
W = _rc.step_out("04_scoring")  # release: was the agent work directory
OUT = f"{W}/results/clinician"; os.makedirs(OUT, exist_ok=True)
proto = yaml.safe_load(open("shared/canonical/protocol_v1.4.yaml"))
rub = proto["scoring"]["llm_judge"]["rubric"]
from prometheus_eval.prompts import SCORE_RUBRIC_TEMPLATE  # noqa
RUBRIC_TEXT = SCORE_RUBRIC_TEMPLATE.format(**rub)

js = pd.read_parquet(f"{W}/results/judge_scores.parquet")
js = js[js.judge_called].copy()
js["family"] = np.where(js.reader.str.startswith("qwen"), "qwen", "extractive")
ds = pd.read_csv("data/dataset.csv")[["person_id", "qa_index", "question", "answer"]]
pairs = pd.read_parquet(f"{W}/results/inputs/pairs_v14.parquet", columns=["person_id", "qa_index", "reader", "condition", "prediction"])
js = js.merge(ds, on=["person_id", "qa_index"]).merge(pairs, on=["person_id", "qa_index", "reader", "condition"])
js = js.sort_values(["reader", "condition", "person_id", "qa_index"]).reset_index(drop=True)
rng = np.random.default_rng(42)
picked = []
for s in range(1, 6):
    taken = {}
    for fam in ["extractive", "qwen"]:
        pool = js[(js.score == s) & (js.family == fam)]
        k = min(10, len(pool)); sel = pool.iloc[rng.choice(len(pool), size=k, replace=False)] if k else pool.iloc[:0]
        taken[fam] = sel
    for fam, other in [("extractive", "qwen"), ("qwen", "extractive")]:
        short = 10 - len(taken[fam])
        if short > 0:
            pool = js[(js.score == s) & (js.family == other) & (~js.index.isin(taken[other].index))]
            k = min(short, len(pool))
            if k:
                taken[other] = pd.concat([taken[other], pool.iloc[rng.choice(len(pool), size=k, replace=False)]])
    picked += [taken["extractive"], taken["qwen"]]
S = pd.concat(picked)
assert not S.duplicated(["person_id", "qa_index", "reader", "condition"]).any()
S = S.iloc[rng.permutation(len(S))].reset_index(drop=True)          # shuffle so order reveals nothing
S["item_id"] = [f"CV{i + 1:03d}" for i in range(len(S))]
key = S[["item_id", "person_id", "qa_index", "reader", "condition", "family", "score"]].rename(columns={"score": "judge_score"})
key.to_csv(f"{OUT}/clinician_validation_100_key.csv", index=False)

wb = Workbook(); ws = wb.active; ws.title = "items"
hdr = ["item_id", "question", "gold_answer", "candidate_answer", "rubric", "clinician1_score", "clinician2_score"]
ws.append(hdr)
for _, r in S.iterrows():
    ws.append([r.item_id, r.question, r.answer, r.prediction.strip(), RUBRIC_TEXT, None, None])
for c, wdt in zip("ABCDEFG", [9, 60, 60, 60, 60, 16, 16]):
    ws.column_dimensions[c].width = wdt
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for cell in ws[1]:
    cell.font = Font(bold=True)
ins = wb.create_sheet("instructions")
for line in ["Score each candidate answer against the gold answer with the rubric (integer 1-5) in clinician1_score / clinician2_score.",
             "Grade only agreement with the gold answer (clinical notes are not provided, same as the judge). Do not consult other raters.",
             "Rubric:", RUBRIC_TEXT]:
    ins.append([line])
ins.column_dimensions["A"].width = 150
wb.save(f"{OUT}/clinician_validation_100.xlsx")
print(len(S), key.groupby(["judge_score", "family"]).size().to_dict())
